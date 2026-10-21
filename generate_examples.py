"""검증용 예제 파일 생성; 실제 논문 연구 결과나 원문을 포함하지 않습니다."""
from pathlib import Path
from io import BytesIO
import fitz
import pandas as pd
from openpyxl import load_workbook

from prompt_templates import ASSIGNMENT, READING, fields_for
from pdf_analysis import PdfDocument, PageText, build_evidence_packets, coverage
from research_ai import normalize_analysis, analysis_summary
from reporting import create_pdf_report, create_excel

OUT=Path(__file__).parent / "examples"
OUT.mkdir(exist_ok=True)
paper={"자료명/저자사항":"(가상자료) 데이터베이스 분석 연구 / 홍길동",
       "발행년":"2026", "발행자":"가상대학교 대학원", "DB":"학위논문", "controlno":"DEMO-TEST"}
data=pd.DataFrame([paper])
metadata={"school":"가상대학교", "department":"연구방법론 전공", "course":"AI 연구분석 실습",
          "student":"예시학생", "professor":"", "date":"2026-10-21"}
doc_pages=[]
for i in range(1,106):
    if i==73:
        text="RESULTS\nThis is a fictional example result of a synthetic study. No measured effect is asserted."
    elif i==102:
        text="CONCLUSION\nThis is a synthetic conclusion solely for UI demonstration."
    else:
        text=f"The synthetic document contains illustrative methods at PDF page {i}. " * 4
    doc_pages.append(PageText(i,text))
doc=PdfDocument(105,tuple(doc_pages))
raw7={"items":[{"key":key,"text": "가상의 연구 결과로 실제 논문 사실이 아닙니다.",
               "evidence":[{"page":73,"quote":"This is a fictional example result of a synthetic study."}]
               if key=="주요연구결과" else []} for key in fields_for(READING)]}
result7=normalize_analysis(raw7,READING,doc)
result7.update({"model":"gpt-4.1-mini","ai_calls":1,"strategy":"실습용 가상 응답",
                "coverage":coverage(doc,build_evidence_packets(doc))})
raw6={"items":[{"key":key,"text":"가상 서식 테스트를 위한 문장입니다.","evidence":[]} for key in fields_for(ASSIGNMENT)]}
result6=normalize_analysis(raw6,ASSIGNMENT,doc)
result6.update({"model":"gpt-4.1-mini","ai_calls":1,"strategy":"실습용 가상 응답",
                "coverage":coverage(doc,build_evidence_packets(doc))})
years=pd.DataFrame([{"발행연도":"2026","건수":1}])
types=pd.DataFrame([{"자료유형":"학위논문","건수":1}])
for mode,result,filename in [
  (READING,result7,"v2_2_7문장_연구노트_예시.pdf"),
  (ASSIGNMENT,result6,"v2_2_6항목_과제보고서_예시.pdf")]:
    note="실습 예시: 이 부분은 사용자가 직접 작성한 연구 질문을 넣는 곳입니다." if mode==READING else ""
    pdf_bytes=create_pdf_report("가상 데이터베이스 연구",metadata,paper,analysis_summary(result),years,types,
                 "실습용 가상자료",summary_mode=mode,analysis_result=result,research_note=note)
    path=OUT/filename
    path.write_bytes(pdf_bytes)
    with fitz.open(stream=pdf_bytes,filetype="pdf") as f:
        print(filename,'pages',len(f),'bytes',len(pdf_bytes))
        text='\n'.join(p.get_text() for p in f)
        assert "AI 논문 요약" in text and "자료 유형별 참고문헌" in text
        if mode==READING:
            assert "7문장" in text and "연구 질문" in text and "PDF 73쪽" in text
    xlsx_bytes=create_excel(data,"가상 데이터베이스 연구",analysis_summary(result),paper,
                             summary_mode=mode,analysis_result=result,research_note=note)
    xlsx=OUT/("v2_2_7문장_연구노트_예시.xlsx" if mode==READING else "v2_2_6항목_과제보고서_예시.xlsx")
    xlsx.write_bytes(xlsx_bytes)
    wb=load_workbook(BytesIO(xlsx_bytes))
    print(xlsx.name,wb.sheetnames)
