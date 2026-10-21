"""원우별 연구 결과를 PDF/Excel 파일(메모리)로 작성."""

from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (HRFlowable, KeepTogether, PageBreak, Paragraph,
                                SimpleDocTemplate, Spacer, Table, TableStyle)

from paperlab import (API_NAME, SOURCE_URL, SUMMARY_FIELDS,
                      calculate_trends, make_reference)
from prompt_templates import ASSIGNMENT, READING, fields_for, DISPLAY_NAMES, mode_label
from citation_styles import (REFERENCE_TEMPLATES, THESIS, THESIS_EXAMPLE,
                             complete_citation, format_reference, missing_citation_fields)


def _excel_safe(value):
    """외부 데이터의 셀 수식 실행 방지 + Excel 셀 길이 제한."""
    if isinstance(value, str):
        value = value[:32000]
        if value.lstrip().startswith(("=", "+", "-", "@")):
            return "'" + value
    return value


def create_excel(df: pd.DataFrame, keyword: str, summary: dict | None = None,
                 chosen_paper: dict | None = None,
                 data_origin: str = "국회도서관 API 검색 서지정보",
                 citation_record: dict | None = None,
                 summary_mode: str = ASSIGNMENT,
                 analysis_result: dict | None = None,
                 research_note: str = "") -> bytes:
    """과제 부록용 다중 시트 Excel. 6항목/7문장·근거·연구노트 출력."""
    year_df, type_df = calculate_trends(df)
    # API 서지정보만으로는 모든 자료의 정식 참고문헌을 만들 수 없습니다.
    # 선택한 자료에 사용자가 입력/검증한 메타데이터만 덮어씁니다.
    selected_id = str((chosen_paper or {}).get("controlno", "")).strip()
    selected_title = str((chosen_paper or {}).get("자료명/저자사항", "")).strip()
    reference_rows = []
    for i, row in enumerate(df.to_dict("records"), 1):
        row_citation = complete_citation(row)
        is_selected = bool(chosen_paper) and (
            (selected_id and str(row.get("controlno", "")).strip() == selected_id)
            or (not selected_id and selected_title
                and str(row.get("자료명/저자사항", "")).strip() == selected_title)
        )
        if is_selected and citation_record:
            row_citation = complete_citation(row, citation_record)
        reference_rows.append({
            "번호": i,
            "자료유형": row_citation["type"],
            "참고문헌(자료유형별)": format_reference(row_citation),
            "미확인 항목": ", ".join(missing_citation_fields(row_citation)) or "없음 (원문 대조 권장)",
            "관리번호": row.get("controlno", ""), "출처": data_origin,
        })
    references = pd.DataFrame(reference_rows)
    if references.empty:
        references = pd.DataFrame(columns=["번호", "자료유형", "참고문헌(자료유형별)",
                                           "미확인 항목", "관리번호", "출처"])
    # 과제에 사용할 자료유형별 형식을 Excel에서도 별도 확인할 수 있도록 안내 시트를 추가합니다.
    template_sheet = pd.DataFrame([
        {"자료유형": kind, "표기 기준": template, "예시": THESIS_EXAMPLE if kind == THESIS else ""}
        for kind, template in REFERENCE_TEMPLATES.items()
    ])
    source = pd.DataFrame([
        {"항목": "데이터 출처", "내용": data_origin},
        {"항목": "원천 출처", "내용": SOURCE_URL if "API" in data_origin else "사용자 직접 제공"},
        {"항목": "실행 검색어", "내용": keyword},
        {"항목": "통계 범위", "내용": "수집한 검색 결과만 해당하며 전체 연구 동향이 아님"},
        {"항목": "참고문헌 유의", "내용": "자료 유형별 과제용 표기 형식입니다. 미확인 항목은 원문에서 직접 확인·수정하세요."},
        {"항목": "AI 유의", "내용": "AI 출력은 원문과 대조하여 사실 확인 필요"},
    ])
    active_fields = fields_for(summary_mode)
    summary_rows = []
    if summary:
        summary_rows = [{"항목": DISPLAY_NAMES.get(key, key),
                         "AI 요약": summary.get(key, "원문에서 확인되지 않음")}
                        for key in active_fields]
    if chosen_paper:
        summary_rows.insert(0, {"항목": "선택 논문", "AI 요약": chosen_paper.get("자료명/저자사항", "")})
    sheets = {
        "검색결과": df,
        "연도별통계": year_df,
        "유형별통계": type_df,
        "참고문헌": references,
        "AI요약": pd.DataFrame(summary_rows, columns=["항목", "AI 요약"]),
        "출처_주의사항": source,
        "참고문헌_표기기준": template_sheet,
    }
    if analysis_result:
        evidence_rows = []
        for item in analysis_result.get("items", []):
            for ev in item.get("evidence", []):
                evidence_rows.append({
                    "분석 항목": DISPLAY_NAMES.get(item["key"], item["key"]),
                    "근거 위치": f"PDF 기준 {ev['page']}쪽" if ev["page"] is not None else "사용자 초록",
                    "검증된 원문 구절": ev["quote"],
                    "검증 범위": ev["status"],
                })
        if not evidence_rows:
            evidence_rows = [{"분석 항목": "근거 미확보", "근거 위치": "-",
                              "검증된 원문 구절": "원문 일치가 검증된 직접 인용 없음",
                              "검증 범위": "원문 수동 검증 필요"}]
        sheets["원문근거_대조"] = pd.DataFrame(evidence_rows)
        cov = analysis_result.get("coverage", {})
        meta_rows = [
            ("요약 모드", mode_label(summary_mode)),
            ("AI 모델", analysis_result.get("model", "")),
            ("AI 요청 횟수", analysis_result.get("ai_calls", 0)),
            ("분석 전략", analysis_result.get("strategy", "")),
            ("PDF 전체 페이지", cov.get("total_pages", 0)),
            ("PDF 텍스트 추출 페이지", cov.get("text_pages", 0)),
            ("AI 대표 페이지 수", cov.get("selected_page_count", 0)),
            ("AI 분석 대표 페이지", ", ".join(map(str, cov.get("selected_pages", [])))),
            ("검증된 직접 인용 수", analysis_result.get("evidence_verified_count", 0)),
            ("일치 실패로 제외한 인용 수", analysis_result.get("unverified_quote_count", 0)),
            ("검증 주의", "인용문과 페이지의 일치만 확인. 학술적 의미·결론은 연구자가 원문에서 검토"),
        ]
        sheets["AI분석_범위"] = pd.DataFrame(meta_rows, columns=["항목", "결과"])
    if summary_mode == READING:
        sheets["나의_연구노트"] = pd.DataFrame([
            {"항목": "내 생각 / 추가 연구 질문", "직접 작성한 내용": research_note.strip() or "(사용자 미작성)"},
            {"항목": "작성자", "직접 작성한 내용": "AI가 자동 생성한 내용이 아닙니다."},
        ])
    memory = BytesIO()
    with pd.ExcelWriter(memory, engine="openpyxl") as writer:
        for name, data in sheets.items():
            if data.empty and len(data.columns) == 0:
                data = pd.DataFrame({"안내": ["수집 자료가 없습니다."]})
            data = data.map(_excel_safe) if hasattr(data, "map") else data.applymap(_excel_safe)
            data.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.sheet_view.showGridLines = False
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            ws.row_dimensions[1].height = 27
            for cell in ws[1]:
                cell.font = Font(name="맑은 고딕", bold=True, color="FFFFFF", size=11)
                cell.fill = PatternFill("solid", fgColor="17365D")
                cell.alignment = Alignment(vertical="center")
            for col_idx, col_cells in enumerate(ws.columns, 1):
                letter = get_column_letter(col_idx)
                col_name = str(ws.cell(1, col_idx).value or "")
                target_width = 66 if col_name in ["자료명/저자사항", "참고문헌(자료유형별)", "표기 기준", "예시", "검증된 원문 구절", "직접 작성한 내용"] else (55 if col_name in ["AI 요약", "내용", "결과"] else 24)
                ws.column_dimensions[letter].width = target_width
                for cell in list(col_cells)[1:]:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    cell.font = Font(name="맑은 고딕", size=10, color="16724F")
                    if cell.row % 2 == 0:
                        cell.fill = PatternFill("solid", fgColor="F3F7FB")
            # 많은 논문에 긴 제목이 있을 수 있으므로 행 높이 과도 확장 방지
            for r in range(2, ws.max_row + 1):
                tall = {"참고문헌": 58, "참고문헌_표기기준": 58,
                        "AI요약": 64, "원문근거_대조": 72,
                        "AI분석_범위": 48, "나의_연구노트": 76}
                ws.row_dimensions[r].height = tall.get(name, 31)
    return memory.getvalue()


def _register_korean_font() -> tuple[str, str]:
    """Streamlit Linux packages.txt의 fonts-nanum, Windows 맑은 고딕 지원."""
    for regular, bold in [
        ("/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
         "/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf"),
        ("C:/Windows/Fonts/malgun.ttf", "C:/Windows/Fonts/malgunbd.ttf"),
        ("/usr/share/fonts/truetype/nanum/NanumBarunGothic.ttf",
         "/usr/share/fonts/truetype/nanum/NanumBarunGothicBold.ttf"),
    ]:
        if Path(regular).exists():
            if "PaperKorean" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("PaperKorean", regular))
                pdfmetrics.registerFont(TTFont("PaperKoreanBold", bold if Path(bold).exists() else regular))
            return "PaperKorean", "PaperKoreanBold"
    # 임베디드 한글 폰트를 찾지 못한 경우 CID 폰트로 대체
    if "HYGothic-Medium" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont("HYGothic-Medium"))
    return "HYGothic-Medium", "HYGothic-Medium"


def _para(text: object, style: ParagraphStyle) -> Paragraph:
    value = escape(str(text if text is not None else ""))
    value = value.replace("\n", "<br/>")
    return Paragraph(value, style)


def create_pdf_report(
    keyword: str,
    metadata: dict,
    selected_paper: dict,
    summary: dict,
    stats_df: pd.DataFrame,
    type_df: pd.DataFrame,
    evidence_source: str,
    data_origin: str = "국회도서관 API 검색 서지정보",
    citation_record: dict | None = None,
    summary_mode: str = ASSIGNMENT,
    analysis_result: dict | None = None,
    research_note: str = "",
) -> bytes:
    """과제 6항목 또는 읽기용 7문장 + 페이지별 인용 근거 보고서(PDF)."""
    regular, bold = _register_korean_font()
    ink = colors.HexColor("#20364E")
    muted = colors.HexColor("#506478")
    accent = colors.HexColor("#1E7186")
    light = colors.HexColor("#ECF5F7")
    doc_buffer = BytesIO()
    doc = SimpleDocTemplate(doc_buffer, pagesize=A4,
                            leftMargin=22 * mm, rightMargin=22 * mm,
                            topMargin=20 * mm, bottomMargin=20 * mm,
                            title="논문 검색 및 AI 분석 보고서")
    styles = {
        "title": ParagraphStyle("report_title", fontName=bold, fontSize=22, leading=34,
                                textColor=ink, alignment=TA_CENTER, spaceAfter=10),
        "subtitle": ParagraphStyle("report_subtitle", fontName=regular, fontSize=11,
                                   leading=18, textColor=muted, alignment=TA_CENTER),
        "section": ParagraphStyle("report_section", fontName=bold, fontSize=13,
                                  leading=20, textColor=accent, spaceBefore=8, spaceAfter=6),
        "body": ParagraphStyle("report_body", fontName=regular, fontSize=9.6,
                               leading=15.5, textColor=ink, spaceAfter=4, wordWrap="CJK"),
        "small": ParagraphStyle("report_small", fontName=regular, fontSize=8.5,
                                leading=14, textColor=muted, wordWrap="CJK"),
        "label": ParagraphStyle("report_label", fontName=bold, fontSize=9.4,
                                leading=16, textColor=ink, wordWrap="CJK"),
    }
    body = []
    body.append(Spacer(1, 38 * mm))
    cover_title = "연구자료 검색 · 분석 과제" if summary_mode == ASSIGNMENT else "논문 읽기 · 연구노트"
    body.append(_para(cover_title, styles["title"]))
    body.append(_para(f"서지정보 및 {mode_label(summary_mode)} 분석 보고서", styles["subtitle"]))
    body.append(Spacer(1, 14 * mm))
    body.append(HRFlowable(width="100%", thickness=1.5, color=accent))
    body.append(Spacer(1, 18 * mm))
    form_items = [
        ("학교", metadata.get("school", "")),
        ("대학원 / 전공", metadata.get("department", "")),
        ("교과목", metadata.get("course", "")),
        ("담당교수", metadata.get("professor", "")),
        ("제출자", metadata.get("student", "")),
        ("제출일", metadata.get("date", date.today().isoformat())),
        ("검색 주제", keyword),
    ]
    table_data = [[_para(k, styles["label"]), _para(v or "(미입력)", styles["body"])]
                  for k, v in form_items]
    cover_table = Table(table_data, colWidths=[40 * mm, 111 * mm], hAlign="CENTER")
    cover_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F0F4F8")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 11),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 11),
        ("LEFTPADDING", (0, 0), (-1, -1), 10),
        ("LINEBELOW", (0, 0), (-1, -1), 0.45, colors.HexColor("#DCE5ED")),
    ]))
    body.append(cover_table)
    body.append(PageBreak())

    body.append(_para("1. 연구 및 검색 개요", styles["section"]))
    body.append(_para(f"검색어: {keyword}", styles["body"]))
    body.append(_para(f"자료 출처: {data_origin}" + (f" ({SOURCE_URL})" if "API" in data_origin else ""), styles["body"]))
    body.append(_para("통계는 이번에 수집한 검색 결과에만 해당합니다. 전체 학술 연구 동향으로 일반화할 수 없습니다.", styles["small"]))
    body.append(_para("2. 선택한 연구자료 서지정보", styles["section"]))
    for field, label in [("자료명/저자사항", "자료명/저자사항"),
                         ("발행자", "발행자"), ("발행년", "발행연도"),
                         ("DB", "자료유형"), ("controlno", "관리번호")]:
        body.append(_para(f"{label}: {selected_paper.get(field, '확인 불가')}", styles["body"]))

    body.append(_para(f"3. AI 논문 요약 ({mode_label(summary_mode)})", styles["section"]))
    rows = []
    for i, key in enumerate(fields_for(summary_mode), 1):
        rows.append([_para(f"{i}. {DISPLAY_NAMES.get(key, key)}", styles["label"]),
                     _para(summary.get(key, "원문에서 확인되지 않음"), styles["body"])])
    sum_table = Table(rows, colWidths=[48 * mm, 109 * mm], repeatRows=0, hAlign="LEFT")
    sum_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BACKGROUND", (0, 0), (0, -1), light),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 9),
        ("LINEBELOW", (0, 0), (-1, -1), 0.35, colors.HexColor("#DBE7ED")),
    ]))
    body.append(sum_table)
    body.append(Spacer(1, 6 * mm))
    body.append(_para("분석 근거: " + evidence_source, styles["small"]))

    if analysis_result:
        info = analysis_result.get("coverage", {})
        body.append(_para("원문 근거·검증 범위", styles["section"]))
        if info.get("total_pages", 0):
            body.append(_para(
                f"PDF 전체 {info.get('total_pages', 0)}쪽에서 텍스트 추출 "
                f"({info.get('text_pages', 0)}쪽 유효); AI 근거에 포함한 대표 페이지 "
                f"{info.get('selected_page_count', 0)}쪽. "
                "전 페이지를 AI에 그대로 전송하거나 완독한 것은 아닙니다.", styles["small"]))
        evidence_rows = []
        for item in analysis_result.get("items", []):
            for quote in item.get("evidence", []):
                page_label = f"PDF {quote['page']}쪽" if quote.get("page") is not None else "초록"
                evidence_rows.append([
                    _para(DISPLAY_NAMES.get(item["key"], item["key"]), styles["label"]),
                    _para(f"{page_label} · {quote['quote'][:210]}", styles["small"]),
                ])
        if evidence_rows:
            evidence_table = Table(evidence_rows, colWidths=[36 * mm, 121 * mm],
                                   repeatRows=0, hAlign="LEFT")
            evidence_table.setStyle(TableStyle([
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LINEBELOW", (0, 0), (-1, -1), 0.35, colors.HexColor("#DBE7ED")),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            body.append(evidence_table)
        else:
            body.append(_para("원문 문구 일치가 검증된 직접 인용이 없습니다. 원문을 직접 확인해 주세요.", styles["small"]))
        body.append(_para(
            "근거는 PDF 페이지/텍스트 구절의 존재 여부만 기계적으로 검증합니다. "
            "의미 해석의 정확성을 입증하지 않으므로 반드시 원문과 대조하세요.", styles["small"]))

    if summary_mode == READING:
        body.append(_para("나의 생각 / 추가 연구 질문 (사용자 직접 작성)", styles["section"]))
        body.append(_para(research_note[:4000] if research_note.strip() else "(아직 작성하지 않음)", styles["body"]))
        body.append(_para("이 부분은 AI가 작성한 내용이 아닙니다.", styles["small"]))

    stat_block = [_para("4. 수집 자료 통계", styles["section"])]
    if not stats_df.empty:
        text = ", ".join(f"{r['발행연도']}년 {r['건수']}건" for r in stats_df.to_dict("records"))
        stat_block.append(_para("발행연도별: " + text[:1000], styles["body"]))
    if not type_df.empty:
        text = ", ".join(f"{r['자료유형']} {r['건수']}건" for r in type_df.to_dict("records"))
        stat_block.append(_para("자료유형별: " + text[:1000], styles["body"]))
    body.append(KeepTogether(stat_block))
    body.append(_para("5. 참고문헌 및 검증 유의사항", styles["section"]))
    citation = complete_citation(selected_paper, citation_record)
    body.append(_para(f"자료유형: {citation['type']}", styles["label"]))
    body.append(_para(format_reference(citation), styles["body"]))
    missing = missing_citation_fields(citation)
    if missing:
        body.append(_para("제출 전 확인할 항목: " + ", ".join(missing), styles["small"]))
    body.append(_para(
        "위 참고문헌은 사용자 지정 과제용 표기 양식입니다. API에 없는 정보는 '[미확인]'으로 표시하며, "
        "정확한 저자·기관·권호·쪽수·DOI는 원문을 대조하고 작성자가 확인해야 합니다.", styles["small"]))

    body.append(Spacer(1, 3 * mm))
    body.append(_para(
        "본 요약은 AI 생성 문장입니다. 원문 전체의 정확성·일관성을 보장하지 않습니다. "
        "표절과 저작권 규정을 준수하고 각 주장·근거를 확인하세요.", styles["small"]))

    # 부록 제목과 표/예시를 한 덩어리로 묶어 제목 고립과 불필요한 빈 페이지 방지.
    appendix = [
        _para("부록. 자료 유형별 참고문헌 표기 기준", styles["section"]),
        _para("아래 3가지 과제용 표기 형식 중 자료의 실제 유형에 맞게 선택하여 작성하세요.", styles["small"]),
    ]
    guide_rows = [[_para("자료 유형", styles["label"]), _para("과제용 표기 형식", styles["label"])]]
    for kind, template in REFERENCE_TEMPLATES.items():
        guide_rows.append([_para(kind, styles["label"]), _para(template, styles["small"])])
    guide_table = Table(guide_rows, colWidths=[36 * mm, 121 * mm], repeatRows=1, hAlign="LEFT")
    guide_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BACKGROUND", (0, 0), (-1, 0), light),
        ("LINEBELOW", (0, 0), (-1, 0), 0.6, accent),
        ("LINEBELOW", (0, 1), (-1, -1), 0.3, colors.HexColor("#DBE7ED")),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
    ]))
    appendix.append(guide_table)
    appendix.append(Spacer(1, 2 * mm))
    appendix.append(_para("학위논문 표기 예시: " + THESIS_EXAMPLE, styles["small"]))
    body.append(KeepTogether(appendix))

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont(regular, 8)
        canvas.setFillColor(muted)
        canvas.drawRightString(A4[0] - 22 * mm, 12 * mm, f"{document.page}")
        canvas.restoreState()

    doc.build(body, onFirstPage=footer, onLaterPages=footer)
    return doc_buffer.getvalue()
