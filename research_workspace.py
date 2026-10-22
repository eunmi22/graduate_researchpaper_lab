"""v2.3 연구 작업공간: 메뉴 간 자료 연결, 연구 기록 및 다운로드.

OpenAI 키와 PDF 원문 bytes는 작업공간 JSON/Excel에 포함하지 않습니다.
AI가 제안한 연구 질문은 선행연구 검증 전까지 '후보'입니다.
"""
from __future__ import annotations

import hashlib
import io
import json
from datetime import date
from typing import Mapping

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from citation_styles import complete_citation, format_reference
from prompt_templates import ASSIGNMENT, READING
from research_ai import analysis_summary

DIMENSIONS = (
    '논문 주요 내용', '활용 기술', '연구 방법', '연구 결과',
    '연구 기여도', '연구 한계', '새로운 쟁점',
)
INCOMPLETE = '원문 확인 및 사용자 검토 필요'
PROJECT_VERSION = '2.3'


def paper_uid(row: Mapping, digest: str = '') -> str:
    """동일 제목이라도 원문이 다른 경우를 구분하도록 자료 ID에 PDF 해시 포함."""
    label = str(row.get('controlno') or row.get('자료명/저자사항') or '')
    return hashlib.sha256((label + '|' + digest).encode('utf-8')).hexdigest()[:20]


def source_dimensions(result: dict | None, note: str = '') -> dict[str, str]:
    """AI 6/7 요약에 실제 존재하는 내용만 7개 연구 항목에 대응."""
    out = {key: INCOMPLETE for key in DIMENSIONS}
    if not result:
        if note.strip():
            out['새로운 쟁점'] = note.strip()[:2000] + ' (사용자 생각, 미검증)'
        return out
    data = analysis_summary(result)
    if result.get('mode') == READING:
        out.update({
            '논문 주요 내용': ' / '.join(filter(None, [data.get('연구배경'), data.get('연구질문')])),
            '활용 기술': data.get('연구방법·핵심기술', INCOMPLETE),
            '연구 방법': data.get('연구방법·핵심기술', INCOMPLETE),
            '연구 결과': data.get('주요연구결과', INCOMPLETE),
            '연구 기여도': data.get('연구의의·활용', INCOMPLETE),
            '연구 한계': data.get('연구한계·후속연구', INCOMPLETE),
        })
    elif result.get('mode') == ASSIGNMENT:
        out.update({
            '논문 주요 내용': ' / '.join(filter(None, [data.get('주제'), data.get('목적')])),
            '활용 기술': data.get('다룬기술', INCOMPLETE),
            '연구 기여도': data.get('활용분야', INCOMPLETE) + ' (활용 분야 요약에서 가져옴; 기여도 직접 검증 필요)',
            '연구 한계': data.get('주요한계', INCOMPLETE),
        })
    if note.strip():
        out['새로운 쟁점'] = note.strip()[:2000] + ' (사용자 생각, 미검증)'
    return {key: str(value or INCOMPLETE)[:2200] for key, value in out.items()}


def build_entry(paper: Mapping, citation_edits: Mapping | None,
                analysis_result: dict | None, note: str,
                digest: str = '', origin: str = '') -> dict:
    """서지정보·AI 요약·사용자 생각을 구분하여 라이브러리 레코드 생성."""
    citation = complete_citation(paper, citation_edits)
    return {
        'id': paper_uid(paper, digest),
        'title': str(citation.get('title') or paper.get('자료명/저자사항') or '')[:500],
        'reference': format_reference(citation),
        'citation': citation,
        'data_origin': str(origin)[:160],
        'saved_on': date.today().isoformat(),
        'notes': str(note or '')[:4000],
        'dimensions': source_dimensions(analysis_result, note),
        'review_status': '개별 논문 요약 / 추가 선행연구 검증 필요',
    }


def literature_rows(library: list[dict]) -> list[dict]:
    rows = []
    for item in library:
        row = {
            '정식 참고문헌': str(item.get('reference', '')),
            '검토 상태': str(item.get('review_status', '미검증')),
        }
        dims = item.get('dimensions', {})
        for key in DIMENSIONS:
            row[key] = str(dims.get(key, INCOMPLETE))
        row['나의 생각·의문'] = str(item.get('notes', ''))
        rows.append(row)
    return rows


def _safe_excel_text(raw: object) -> object:
    """=+-@ 등으로 시작하는 텍스트가 엑셀 수식으로 해석되지 않도록 함."""
    if not isinstance(raw, str):
        return raw
    value = raw.replace('\x00', '')[:32000]
    return "'" + value if value.lstrip().startswith(('=', '+', '-', '@')) else value


def literature_excel(library: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = '선행연구_비교표'
    rows = literature_rows(library)
    cols = ['정식 참고문헌', '검토 상태', *DIMENSIONS, '나의 생각·의문']
    ws.append(cols)
    for item in rows:
        ws.append([_safe_excel_text(item.get(col, '')) for col in cols])
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='17365D')
        cell.font = Font(name='맑은 고딕', bold=True, color='FFFFFF')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 32
    for idx in range(1, len(cols) + 1):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(idx)].width = 65 if idx == 1 else 38
    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height = 96
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if cell.column == len(cols):
                cell.font = Font(name='맑은 고딕', size=10, color='2154A2')  # 사용자 직접 작성
            elif cell.column == 2:
                cell.font = Font(name='맑은 고딕', size=10, color='9B5D13')  # 검토 필요
            else:
                cell.font = Font(name='맑은 고딕', size=10, color='087047')  # 문헌·AI 수집
    ws2 = wb.create_sheet('이용안내')
    ws2.append(['설명', '내용'])
    ws2.append(['정식 참고문헌', '모든 출처는 자료 유형별 정식 표기를 사용합니다. 미확인 메타데이터는 원문 확인 후 수정합니다.'])
    ws2.append(['연구 공백', '나의 의문은 연구 공백으로 확정되지 않았습니다. 추가 선행연구 비교와 지도교수 확인이 필요합니다.'])
    ws2.append(['AI 결과', 'AI 요약의 원문 일치 여부와 내용의 사실 여부는 별도로 검토해야 합니다.'])
    result = io.BytesIO()
    wb.save(result)
    return result.getvalue()


def direction_prompt(library: list[dict], research_interest: str,
                     purpose: str = '석사학위논문',
                     available_data: str = '', timeframe: str = '') -> str:
    """사용자가 명시한 논문과 생각에서만 출발하도록 강제하는 연구 검토용 프롬프트."""
    parts = [
        '당신은 대학원 연구 설계 보조자입니다. 아래 자료만을 직접 확인한 사실로 사용하세요.',
        '확인하지 못한 최신 선행연구, 실험 결과, DOI, 신기술 사실을 만들어내지 마세요.',
        '논문에서 확인된 한계 / 사용자가 작성한 의문 / AI의 신규 제안을 각각 구분하세요.',
        '내 연구 주제 후보 3개와 각 RQ, 필요한 자료, 평가 방법, 예상 기여, 현실적 위험, 선행연구 확인용 검색어를 제시하세요.',
        '연구 공백은 현재는 검증 전 후보로 표시하고, 새로운 선행연구 조사가 필요하다고 명시하세요.',
        '참고문헌을 언급할 때는 아래 정식 표기를 그대로 사용하며 약식으로 바꾸지 마세요.',
        f'연구 목적: {purpose[:120]}',
        f'관심 주제: {research_interest[:300] or "[사용자 입력 필요]"}',
        f'확보 가능 자료·도구: {available_data[:800] or "[확인 필요]"}',
        f'연구 기간: {timeframe[:120] or "[확인 필요]"}',
        '\n[검토한 논문과 개인 연구노트]',
    ]
    if not library:
        parts.append('선행연구 비교 데이터가 없습니다. 연구 공백이나 독창성을 확정하지 말고 우선 조사 계획을 제시하세요.')
    for i, item in enumerate(library[:12], 1):
        dims = item.get('dimensions', {})
        parts.append(f"\n{i}. 정식 참고문헌: {item.get('reference', '[자료 확인 필요]')[:1000]}")
        for key in DIMENSIONS:
            parts.append(f"- {key}: {str(dims.get(key, INCOMPLETE))[:600]}")
        parts.append(f"- 사용자가 작성한 의문: {str(item.get('notes', ''))[:900]}")
    parts.append('\n[출력 형식] 주제 후보 1~3 / 각 후보의 RQ·필요 자료·방법·평가 지표·예상 기여·검증할 문헌 / 지도교수 상담 질문')
    return '\n'.join(parts)


def export_snapshot(library: list[dict], workspace: Mapping | None = None) -> bytes:
    """개인 키·원본 PDF가 없는 JSON만 다운로드; 크기 제한으로 민감 내용 과도한 누적 방지."""
    payload = {
        'schema': 'graduate-paper-lab-workspace', 'version': PROJECT_VERSION,
        'export_date': date.today().isoformat(),
        'library': library[:100],
        'workspace': {str(k)[:80]: v for k, v in (workspace or {}).items()
                      if k not in {'openai_key', 'pdf_bytes', 'pdf_doc'}},
    }
    raw = json.dumps(payload, ensure_ascii=False, indent=2).encode('utf-8')
    if len(raw) > 2 * 1024 * 1024:
        raise ValueError('연구 프로젝트 내보내기 크기는 2MB 이하여야 합니다.')
    return raw


def import_snapshot(raw: bytes) -> tuple[list[dict], dict]:
    if not raw or len(raw) > 2 * 1024 * 1024:
        raise ValueError('2MB 이하의 프로젝트 JSON을 사용하세요.')
    try:
        value = json.loads(raw.decode('utf-8-sig'))
    except (ValueError, UnicodeDecodeError) as exc:
        raise ValueError('읽을 수 있는 UTF-8 JSON 형식이 아닙니다.') from exc
    if not isinstance(value, dict) or value.get('schema') != 'graduate-paper-lab-workspace':
        raise ValueError('대학원 AI 논문 연구 도우미의 프로젝트 파일이 아닙니다.')
    lib = value.get('library')
    info = value.get('workspace', {})
    if not isinstance(lib, list) or len(lib) > 100 or not isinstance(info, dict):
        raise ValueError('프로젝트 데이터 구조가 올바르지 않습니다.')
    sanitized = []
    for item in lib:
        if not isinstance(item, dict) or not isinstance(item.get('citation'), dict):
            raise ValueError('선행연구 기록 구조가 올바르지 않습니다.')
        record = {str(k)[:80]: v for k, v in item.items() if k in {
            'id', 'title', 'citation', 'data_origin', 'saved_on', 'notes',
            'dimensions', 'review_status'
        }}
        record['reference'] = format_reference(record['citation'])
        record['notes'] = str(record.get('notes', ''))[:4000]
        dims = record.get('dimensions', {})
        if not isinstance(dims, dict):
            dims = {}
        record['dimensions'] = {name: str(dims.get(name, INCOMPLETE))[:2200] for name in DIMENSIONS}
        sanitized.append(record)
    return sanitized, {str(k): v for k, v in info.items() if k in {
        'research_interest', 'purpose', 'available_data', 'timeframe',
        'research_plan', 'methods_note', 'results_note', 'introduction', 'discussion', 'research_candidates'
    }}
